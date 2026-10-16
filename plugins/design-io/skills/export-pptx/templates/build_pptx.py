#!/usr/bin/env python3
"""Режим editable, шаг 2: slides.json (от extract.mjs) -> PPTX с нативными TextBox.

python-pptx импортируется внутри main(), после разбора аргументов:
иначе `--help` падал бы с ImportError там, где зависимость не поставлена.
"""
import json
import os
import sys
import tempfile
import urllib.request

USAGE = """Usage: build_pptx.py <slides.json> [out.pptx]

  slides.json   вывод extract.mjs
  out.pptx      куда писать (default: то же имя с расширением .pptx)
  -h, --help    эта справка

Зависимости: pip install python-pptx pillow"""

EMU_PER_PX = 9525  # PowerPoint: 1 px @96dpi = 9525 EMU


def hex_to_rgb(RGBColor, s):
    s = s.lstrip('#')
    return RGBColor(int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))


def fetch_image(src, tmp):
    if src.startswith('data:'):
        import base64
        head, b64 = src.split(',', 1)
        ext = 'png' if 'png' in head else 'jpg'
        path = os.path.join(tmp, f"img_{abs(hash(src))}.{ext}")
        with open(path, 'wb') as f:
            f.write(base64.b64decode(b64))
        return path
    if src.startswith('http'):
        path = os.path.join(tmp, f"img_{abs(hash(src))}")
        urllib.request.urlretrieve(src, path)
        return path
    if src.startswith('file://'):
        return src.replace('file://', '')
    return src  # относительный путь


def main():
    argv = sys.argv[1:]
    if not argv or argv[0] in ('-h', '--help'):
        print(USAGE)
        sys.exit(0 if argv else 1)

    src_json = argv[0]
    if not os.path.exists(src_json):
        print(f"✗ нет файла: {src_json}", file=sys.stderr)
        sys.exit(1)

    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    def px(v):
        return Emu(int(round(v * EMU_PER_PX)))

    with open(src_json, encoding='utf-8') as f:
        data = json.load(f)
    out = argv[1] if len(argv) > 1 else src_json.replace('.json', '.pptx')
    W, H = data['width'], data['height']

    prs = Presentation()
    prs.slide_width = px(W)
    prs.slide_height = px(H)
    blank = prs.slide_layouts[6]

    tmp = tempfile.mkdtemp()

    for s in data['slides']:
        slide = prs.slides.add_slide(blank)
        if s.get('bg'):
            try:
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = hex_to_rgb(RGBColor, s['bg'])
            except Exception:
                pass

        for item in s['items']:
            x, y, w, h = px(item['x']), px(item['y']), px(item['w']), px(item['h'])

            if item['kind'] == 'rect':
                shp = slide.shapes.add_shape(1, x, y, w, h)  # MSO_SHAPE.RECTANGLE
                shp.fill.solid()
                if item.get('fill'):
                    shp.fill.fore_color.rgb = hex_to_rgb(RGBColor, item['fill'])
                shp.line.fill.background()

            elif item['kind'] == 'text':
                tb = slide.shapes.add_textbox(x, y, w, h)
                tf = tb.text_frame
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                p = tf.paragraphs[0]
                p.text = item['text']
                if not p.runs:
                    continue
                run = p.runs[0]
                run.font.name = item.get('fontFamily') or 'Helvetica'
                run.font.size = Pt(item.get('fontSize', 16) * 0.75)  # px -> pt
                if item.get('color'):
                    run.font.color.rgb = hex_to_rgb(RGBColor, item['color'])
                fw = str(item.get('fontWeight', '400'))
                run.font.bold = fw in ('bold', '600', '700', '800', '900')
                run.font.italic = bool(item.get('italic'))
                a = item.get('align') or 'left'
                p.alignment = {
                    'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER,
                    'right': PP_ALIGN.RIGHT, 'justify': PP_ALIGN.JUSTIFY,
                }.get(a, PP_ALIGN.LEFT)

            elif item['kind'] == 'image':
                try:
                    p = fetch_image(item['src'], tmp)
                    slide.shapes.add_picture(p, x, y, w, h)
                except Exception as e:
                    # Не глотаем молча: пропавшая картинка должна быть видна в логе,
                    # иначе дек уедет заказчику с дырой, и никто не узнает почему.
                    print("img fail:", item['src'], e, file=sys.stderr)

    prs.save(out)
    print("✓", out)


if __name__ == "__main__":
    main()
